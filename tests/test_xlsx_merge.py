from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from automate_lab.xlsx_merge import merge_xlsx, main


def write(path, rows, sheet="sales"):
    book = Workbook()
    book.active.title = sheet
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()


def values(path):
    book = load_workbook(path)
    try:
        return list(book.active.values)
    finally:
        book.close()


def test_reorder_trace_and_repeat_without_ingesting_output(tmp_path):
    write(tmp_path / "a.xlsx", [["product", "qty"], ["001", 3], [None, None]])
    write(tmp_path / "b.xlsx", [["qty", "product"], [4, "B"]])
    (tmp_path / "~$lock.xlsx").write_bytes(b"not an Excel file")
    output = tmp_path / "merged.xlsx"
    before = (tmp_path / "a.xlsx").read_bytes()
    for _ in range(2):
        summary = merge_xlsx(tmp_path, output)
        assert summary["rows"] == 2
        assert summary["files"] == 2
        assert values(output) == [("product", "qty", "_source_file"),
                                  ("001", 3, "a.xlsx"), ("B", 4, "b.xlsx")]
    assert (tmp_path / "a.xlsx").read_bytes() == before


@pytest.mark.parametrize("bad_rows, message", [
    ([["different"], [1]], "header mismatch"),
    ([["product", "product"], [1, 2]], "duplicate"),
    ([["product", None], [1, 2]], "nonempty"),
    ([["product", "_source_file"], [1, 2]], "reserved"),
    ([["product", "qty"], ["A", "=1+1"]], "formula"),
    ([["product", "qty"], ["A", "#DIV/0!"]], "Excel error"),
])
def test_invalid_input_preserves_existing_output(tmp_path, bad_rows, message):
    write(tmp_path / "a.xlsx", [["product", "qty"], ["A", 1]])
    write(tmp_path / "b.xlsx", bad_rows)
    output = tmp_path / "result.xlsx"
    output.write_bytes(b"previous result")
    with pytest.raises(ValueError, match=message):
        merge_xlsx(tmp_path, output)
    assert output.read_bytes() == b"previous result"


def test_sheet_selection_and_missing_sheet(tmp_path):
    write(tmp_path / "a.xlsx", [["x"], [2]], sheet="target")
    assert merge_xlsx(tmp_path, tmp_path / "out.xlsx", sheet="target")["rows"] == 1
    with pytest.raises(ValueError, match="missing sheet"):
        merge_xlsx(tmp_path, tmp_path / "out.xlsx", sheet="missing")


def test_empty_and_no_data(tmp_path):
    with pytest.raises(ValueError, match="No input"):
        merge_xlsx(tmp_path, tmp_path / "out.xlsx")
    write(tmp_path / "a.xlsx", [["x"]])
    with pytest.raises(ValueError, match="No data"):
        merge_xlsx(tmp_path, tmp_path / "out.xlsx")


def test_cli_and_text_that_looks_like_formula(tmp_path, capsys):
    path = tmp_path / "a.xlsx"
    book = Workbook()
    book.active.append(["x"])
    book.active.append(["=literal"])
    book.active["A2"].data_type = "s"
    book.save(path)
    book.close()
    output = tmp_path / "out.xlsx"
    assert main(["--input", str(tmp_path), "--out", str(output)]) == 0
    assert "files=1 rows=1" in capsys.readouterr().out
    book = load_workbook(output)
    assert book.active["A2"].value == "=literal"
    assert book.active["A2"].data_type == "s"
    book.close()
