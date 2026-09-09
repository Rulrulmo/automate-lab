"""Receipt OCR → LLM structured extract → Excel append.

Blog: https://automate-lab.tistory.com/25
Optional deps: pip install 'automate-lab[receipts]'
"""
from __future__ import annotations

import argparse
import json
import os
import re
import traceback
from pathlib import Path

HEADERS = ["date", "merchant", "amount", "currency", "confidence", "source", "notes"]

SCHEMA = {
    "type": "object",
    "properties": {
        "date": {"type": "string", "description": "YYYY-MM-DD"},
        "merchant": {"type": "string"},
        "amount": {"type": "number"},
        "currency": {"type": "string"},
        "confidence": {"type": "number", "description": "0~1"},
        "notes": {"type": "string"},
    },
    "required": ["date", "merchant", "amount", "currency", "confidence"],
    "additionalProperties": False,
}


def amount_candidates(text: str) -> list[int]:
    found: list[int] = []
    for m in re.finditer(r"(?<!\d)(\d{1,3}(?:[,.]\d{3})+|\d+)(?!\d)", text):
        raw = m.group(1).replace(",", "").replace(".", "")
        if raw.isdigit():
            found.append(int(raw))
    return found


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            import pdfplumber
        except ImportError as exc:
            raise ImportError("pdfplumber required: pip install 'automate-lab[receipts]'") from exc
        parts: list[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
        if len(text) >= 40:
            return text
    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise ImportError(
            "Pillow+pytesseract required: pip install 'automate-lab[receipts]'"
        ) from exc
    img = Image.open(path)
    return pytesseract.image_to_string(img, lang="kor+eng")


def llm_extract(ocr_text: str) -> dict:
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise ImportError("openai required: pip install 'automate-lab[receipts]'") from exc

    client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])
    model = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "영수증 OCR 텍스트에서 결제 정보를 추출한다. "
                    "확실하지 않으면 confidence를 낮추고 notes에 이유를 적는다. "
                    "추측으로 금액을 채우지 않는다. 합계/총액/결제금액 우선."
                ),
            },
            {"role": "user", "content": ocr_text[:8000]},
        ],
        response_format={
            "type": "json_schema",
            "json_schema": {
                "name": "receipt_row",
                "schema": SCHEMA,
                "strict": True,
            },
        },
    )
    return json.loads(resp.choices[0].message.content)


def append_row(row: dict, source: str, *, xlsx: Path, review_below: float = 0.7) -> None:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl required: pip install 'automate-lab[receipts]'") from exc

    if not xlsx.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "receipts"
        ws.append(HEADERS)
        wb.save(xlsx)

    wb = load_workbook(xlsx)
    ws = wb["receipts"]
    values = [
        row.get("date"),
        row.get("merchant"),
        row.get("amount"),
        row.get("currency"),
        row.get("confidence"),
        source,
        row.get("notes", ""),
    ]
    ws.append(values)
    conf = float(row.get("confidence") or 0)
    if conf < review_below or row.get("amount") is None:
        if "review" not in wb.sheetnames:
            wb.create_sheet("review")
            wb["review"].append(HEADERS)
        wb["review"].append(values)
    wb.save(xlsx)


def process_one(path: Path, *, xlsx: Path) -> None:
    text = extract_text(path)
    if not text.strip():
        raise RuntimeError(f"no text: {path}")
    data = llm_extract(text)
    cands = amount_candidates(text)
    if cands and data.get("amount") is not None:
        amt = int(data["amount"]) if float(data["amount"]).is_integer() else data["amount"]
        if isinstance(amt, (int, float)) and int(amt) not in cands and amt not in cands:
            notes = (data.get("notes") or "") + f" | amount candidates={cands[:5]}"
            data["notes"] = notes.strip(" |")
            data["confidence"] = min(float(data.get("confidence") or 1), 0.6)
    append_row(data, source=str(path), xlsx=xlsx)


def process_inbox(inbox: Path, *, xlsx: Path) -> tuple[int, int]:
    done = inbox.parent / "done"
    failed = inbox.parent / "failed"
    done.mkdir(exist_ok=True)
    failed.mkdir(exist_ok=True)
    ok = err = 0
    for p in sorted(inbox.iterdir()):
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".pdf", ".png", ".jpg", ".jpeg", ".webp"}:
            continue
        try:
            process_one(p, xlsx=xlsx)
            p.rename(done / p.name)
            ok += 1
        except Exception:
            (failed / (p.name + ".err.txt")).write_text(
                traceback.format_exc(), encoding="utf-8"
            )
            p.rename(failed / p.name)
            err += 1
    return ok, err


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="OCR+LLM receipts → Excel")
    parser.add_argument("--inbox", type=Path, default=Path("inbox"))
    parser.add_argument("--xlsx", type=Path, default=Path("receipts.xlsx"))
    args = parser.parse_args(argv)
    if not args.inbox.is_dir():
        raise SystemExit(f"inbox not found: {args.inbox}")
    ok, err = process_inbox(args.inbox, xlsx=args.xlsx)
    print(f"ok={ok} failed={err} xlsx={args.xlsx}")
    return 0 if err == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
