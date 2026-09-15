"""Combine matching XLSX tables without changing source workbooks."""
from __future__ import annotations

import argparse
import os
from pathlib import Path
from tempfile import NamedTemporaryFile

SOURCE = "_source_file"


def merge_xlsx(folder: Path, output: Path, *, sheet: str | None = None) -> dict:
    """Merge the first (or named) sheet; reject formulas and mismatched headers.

    The first row must contain unique, nonempty text headers. Column order
    may differ. Values are preserved, not formatting or Excel formulas.
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell

    folder, output = Path(folder).resolve(), Path(output).resolve()
    if not folder.is_dir():
        raise ValueError(f"Input folder does not exist: {folder}")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have an .xlsx extension")
    files = sorted(
        (p for p in folder.iterdir() if p.is_file()
         and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")
         and p.resolve() != output),
        key=lambda p: p.name,
    )
    if not files:
        raise ValueError("No input .xlsx files found")

    # Validate every input before replacing an existing output file.
    records = []
    headers = None
    counts = {}
    for path in files:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is not None and sheet not in book.sheetnames:
                raise ValueError(f"{path.name}: missing sheet {sheet!r}")
            ws = book[sheet] if sheet is not None else book.worksheets[0]
            rows = ws.iter_rows()
            first = next(rows, ())
            names = [cell.value for cell in first]
            if any(cell.data_type == "f" for cell in first):
                raise ValueError(f"{path.name}: formula in header")
            if not names or any(not isinstance(v, str) or not v.strip() for v in names):
                raise ValueError(f"{path.name}: headers must be nonempty text")
            names = [v.strip() for v in names]
            if len(set(names)) != len(names) or SOURCE in names:
                raise ValueError(f"{path.name}: duplicate or reserved header")
            if headers is None:
                headers = names
            elif set(names) != set(headers):
                missing = sorted(set(headers) - set(names))
                extra = sorted(set(names) - set(headers))
                raise ValueError(f"{path.name}: header mismatch; missing={missing}, extra={extra}")
            order = [names.index(name) for name in headers]
            count = 0
            for row in rows:
                if all(cell.value is None for cell in row):
                    continue
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError(f"{path.name}: formula found; export values first")
                if any(cell.data_type == "e" for cell in row):
                    raise ValueError(f"{path.name}: Excel error found; fix source first")
                records.append([row[i].value for i in order] + [path.name])
                count += 1
                if len(records) > 1_048_575:
                    raise ValueError("Too many rows for one Excel sheet")
            counts[path.name] = count
        finally:
            book.close()
    if not records:
        raise ValueError("No data rows found")

    result = Workbook(write_only=True)
    ws = result.create_sheet("merged")
    ws.freeze_panes = "A2"
    for values in [headers + [SOURCE], *records]:
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        ws.append(cells)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temp = Path(handle.name)
        result.save(temp)
        os.replace(temp, output)
    finally:
        result.close()
        if temp is not None:
            temp.unlink(missing_ok=True)
    return {"files": len(files), "rows": len(records), "rows_by_file": counts,
            "output": str(output)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Merge XLSX tables with matching headers")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--sheet", help="Sheet name; default: first sheet of each workbook")
    args = parser.parse_args(argv)
    try:
        result = merge_xlsx(args.input, args.out, sheet=args.sheet)
    except (OSError, ValueError) as exc:
        parser.exit(1, f"error: {exc}\n")
    print(f"files={result['files']} rows={result['rows']}")
    for name, count in result["rows_by_file"].items():
        print(f"  {name}: {count} rows")
    print(f"saved: {result['output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
